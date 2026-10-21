# @Auther  : luoscorn
# @Time    : 2026/10/21
# @File    : score_rubric.py
"""xlsx 评分标准 → 结构化评分项"""

import logging
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from utils.audit_type import AUDIT_SCORE_DOC, AuditType, get_score_doc_path
from utils.oss_client import download_file, get_object_url

logger = logging.getLogger(__name__)

_HEADER_NAMES = {
    "dimension": "一级维度",
    "number": "编号",
    "point": "细分评分点",
    "evidence": "视频识别依据",
    "score": "分值",
}
_REQUIRED_COLUMNS = ("number", "point", "score")


@dataclass(frozen=True)
class RubricItem:
    """单个可独立判定的评分点"""
    dimension: str
    number: str
    point: str
    evidence: str
    score: float

    @property
    def label(self) -> str:
        return f"{self.number} {self.point}"


def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def parse_xlsx_rubric(path: str | Path) -> list[RubricItem]:
    """
    解析 xlsx「评分表」工作表为结构化评分项

    :param path: 本地 xlsx 文件路径
    :return: 按表格顺序排列的评分项
    :raises ValueError: 表头缺失、无有效评分项或编号重复
    """
    wb = load_workbook(str(path), data_only=True, read_only=True)
    try:
        ws = wb["评分表"] if "评分表" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    header_idx = -1
    col_idx: dict[str, int] = {}
    for i, row in enumerate(rows):
        labels = [_cell_text(v) for v in row]
        if "编号" in labels and "细分评分点" in labels:
            header_idx = i
            col_idx = {
                key: labels.index(name)
                for key, name in _HEADER_NAMES.items()
                if name in labels
            }
            break

    if header_idx < 0:
        raise ValueError(f"{path} 未找到评分表表头（编号/细分评分点）")
    missing_cols = [key for key in _REQUIRED_COLUMNS if key not in col_idx]
    if missing_cols:
        raise ValueError(f"{path} 评分表缺少列: {missing_cols}")

    def cell(row, key: str) -> str:
        idx = col_idx.get(key)
        return _cell_text(row[idx]) if idx is not None and idx < len(row) else ""

    items: list[RubricItem] = []
    dimension = ""
    for row in rows[header_idx + 1:]:
        # 一级维度按合并单元格向下填充 使每个评分项都带上所属维度
        dimension = cell(row, "dimension") or dimension
        number = cell(row, "number")
        point = cell(row, "point")
        if not number or not point:
            continue
        try:
            score = float(cell(row, "score").replace("分", ""))
        except ValueError:
            continue
        if score <= 0:
            continue
        items.append(RubricItem(
            dimension=dimension,
            number=number,
            point=point,
            evidence=cell(row, "evidence"),
            score=round(score, 2),
        ))

    if not items:
        raise ValueError(f"{path} 未解析到有效评分项")

    numbers = [item.number for item in items]
    if len(numbers) != len(set(numbers)):
        raise ValueError(f"{path} 评分项编号重复: {numbers}")

    return items


def get_score_rubric(audit_type: AuditType) -> list[RubricItem] | None:
    """
    获取结构化评分项（仅 xlsx 标准）；docx 标准返回 None，沿用文本规则

    每次从 OSS 拉取最新 xlsx，下载失败时回退本地缓存
    """
    oss_key = AUDIT_SCORE_DOC.get(audit_type)
    if not oss_key or not oss_key.lower().endswith(".xlsx"):
        return None

    local_path = Path(get_score_doc_path(audit_type))
    local_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        url = get_object_url(oss_key=oss_key)
        download_file(url=url, save_dir=str(local_path.parent), filename=local_path.name)
    except Exception:
        if not local_path.exists():
            raise
        logger.warning("评分表拉取失败，使用本地缓存: %s", local_path, exc_info=True)

    return parse_xlsx_rubric(local_path)
